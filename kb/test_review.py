"""The agronomist review round trip: sheet -> decisions -> kb_treatment review columns."""

import copy
import csv
import sys
import tempfile
import unittest
from pathlib import Path

KB = Path(__file__).resolve().parent
sys.path.insert(0, str(KB))

try:
    import openpyxl
except ImportError:
    openpyxl = None


@unittest.skipUnless(openpyxl, "needs openpyxl")
class ReviewRoundTripTest(unittest.TestCase):
    def test_filled_sheet_becomes_review_columns(self):
        import make_review_sheet as M
        from build_kb import apply_review
        rows = list(csv.DictReader((KB / "kb_treatment.csv").open(encoding="utf-8")))
        with tempfile.TemporaryDirectory() as tmp:
            sheet, decisions = Path(tmp) / "review.xlsx", Path(tmp) / "decisions.csv"
            M.export(sheet, decisions)
            wb = openpyxl.load_workbook(sheet)
            ws = wb["Treatments"]
            header = [c.value for c in ws[1]]
            col = {h: header.index(h) + 1 for h in header}
            ws.cell(2, col["decision"], "approve")                  # rank left blank -> suggested_rank
            ws.cell(3, col["decision"], "approve"); ws.cell(3, col["rank"], 1)
            ws.cell(4, col["decision"], "reject")
            wb.save(sheet)
            M.import_sheet(sheet, decisions)
            reviewed = copy.deepcopy(rows)
            apply_review(reviewed, decisions)
        status = {(r["label"], r["product"]): (r["review_status"], r["display_rank"]) for r in reviewed}
        get = lambda i: status[(ws.cell(i, col["label"]).value, ws.cell(i, col["product"]).value)]
        self.assertEqual(get(2), ("approved", str(ws.cell(2, col["suggested_rank"]).value)))
        self.assertEqual(get(3), ("approved", "1"))
        self.assertEqual(get(4), ("rejected", ""))
        self.assertEqual(sum(s == "pending" for s, _ in status.values()), len(rows) - 3)

    def test_suggested_rank_puts_biological_and_single_actives_first(self):
        import make_review_sheet as M
        rows = [{"label": "x", "product": "A 5% + B 5% SC", "list": "fungicide", "phi_days": "1"},
                {"label": "x", "product": "C 5% SC", "list": "fungicide", "phi_days": "9"},
                {"label": "x", "product": "Bacillus 1% WP", "list": "bio_fungicide", "phi_days": "0"}]
        r = M.suggested(rows)
        self.assertEqual([r[("x", "Bacillus 1% WP")], r[("x", "C 5% SC")], r[("x", "A 5% + B 5% SC")]], [1, 2, 3])


if __name__ == "__main__":
    unittest.main()
