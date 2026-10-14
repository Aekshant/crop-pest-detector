#!/usr/bin/env python3
"""The agronomist review of kb_treatment.csv, as a spreadsheet.

  uvx --with openpyxl python kb/make_review_sheet.py            # -> kb/review/treatment_review.xlsx
  uvx --with openpyxl python kb/make_review_sheet.py --import   # filled sheet -> kb/review_decisions.csv
  uvx --with pdfplumber python kb/build_kb.py                   # applies the decisions

For each (label, product) the reviewer sets decision = approve / reject and, for
approved rows, rank (1 = the product Laya shows first for that problem). Laya
shows only approved rows, best rank first; a problem with none approved gets the
"ask your KVK" reply. suggested_rank is only a starting point: biological first,
then single active ingredient, then shortest waiting period.
"""

import csv
import sys
from pathlib import Path

KB = Path(__file__).resolve().parent
SHEET = KB / "review" / "treatment_review.xlsx"
DECISIONS = KB / "review_decisions.csv"
COLUMNS = ["label", "product", "dose", "phi_days", "type", "actives", "pest_as_listed", "source_url",
           "suggested_rank", "decision", "rank", "reviewer", "notes"]


def suggested(rows):
    def key(r):
        return (0 if r["list"].startswith("bio") else 1, 0 if "+" not in r["product"] else 1, int(r["phi_days"]), r["product"].lower())
    out = {}
    for label in sorted({r["label"] for r in rows}):
        for i, r in enumerate(sorted((r for r in rows if r["label"] == label), key=key), 1):
            out[(r["label"], r["product"])] = i
    return out


def export(sheet=SHEET, decisions=DECISIONS):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    sys.path.insert(0, str(KB))
    from build_kb import actives_in

    rows = list(csv.DictReader((KB / "kb_treatment.csv").open(encoding="utf-8")))
    old = {(d["label"], d["product"]): d for d in csv.DictReader(decisions.open(encoding="utf-8"))} if decisions.exists() else {}
    ranks = suggested(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Treatments"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5D3A")
    for r in sorted(rows, key=lambda r: (r["label"], ranks[(r["label"], r["product"])])):
        d = old.get((r["label"], r["product"]), {})
        ws.append([r["label"], r["product"], r["dose"], int(r["phi_days"]),
                   "biological" if r["list"].startswith("bio") else "chemical", ", ".join(actives_in(r["product"])),
                   r["pest_as_listed"], r["source_url"], ranks[(r["label"], r["product"])],
                   d.get("decision", ""), d.get("rank", ""), d.get("reviewer", ""), d.get("notes", "")])
    dv = DataValidation(type="list", formula1='"approve,reject"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J2:J{ws.max_row}")
    for col, width in zip("ABCDEFGHIJKLM", [22, 46, 44, 9, 11, 30, 40, 30, 9, 11, 7, 14, 30]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=c.column_letter in "BCG")
        for c in row[9:13]:
            c.fill = PatternFill("solid", fgColor="FFF4CC")  # reviewer columns
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    guide = wb.create_sheet("How to review")
    for line in [
        "FarmLens treatment rows, extracted from the CIB&RC 'Major Uses of Pesticides' lists (as on 30.09.2026).",
        "Every dose and waiting period is copied from the list (source_url opens the page). Laya never computes or changes them.",
        "",
        "For each row, set 'decision' to approve or reject. Laya shows only approved rows to farmers.",
        "For approved rows, set 'rank' within each label: 1 = the product Laya names first. Leave blank to use suggested_rank.",
        "Reject anything you would not recommend to a smallholder tomato farmer (availability, toxicity, resistance risk, etc.).",
        "'suggested_rank' is a mechanical default (biological first, then single active, then shortest waiting period), not advice.",
        "Fill 'reviewer' and use 'notes' for anything the developers should know. Then return the file.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    sheet.parent.mkdir(exist_ok=True)
    wb.save(sheet)
    print(f"{len(rows)} rows -> {sheet}")


def import_sheet(sheet=SHEET, decisions=DECISIONS):
    from openpyxl import load_workbook
    ws = load_workbook(sheet, data_only=True)["Treatments"]
    header = [c.value for c in ws[1]]
    out = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(header, values))
        decision = (r.get("decision") or "").strip().lower()
        if decision not in ("approve", "reject", ""):
            raise SystemExit(f"{r['label']} / {r['product']}: decision must be approve or reject, not {decision!r}")
        rank = r.get("rank") or (r.get("suggested_rank") if decision == "approve" else "")
        out.append({"label": r["label"], "product": r["product"], "decision": decision, "rank": rank or "",
                    "reviewer": r.get("reviewer") or "", "notes": r.get("notes") or ""})
    with decisions.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["label", "product", "decision", "rank", "reviewer", "notes"])
        w.writeheader()
        w.writerows(out)
    print(f"{sum(o['decision'] == 'approve' for o in out)} approved, {sum(o['decision'] == 'reject' for o in out)} rejected, "
          f"{sum(not o['decision'] for o in out)} undecided -> {decisions}")


if __name__ == "__main__":
    import_sheet() if "--import" in sys.argv else export()
